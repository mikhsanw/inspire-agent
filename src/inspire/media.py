"""Indra agent: ubah suara, gambar, file, dan video jadi teks."""

from __future__ import annotations

import asyncio
import base64
import io
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

import httpx
import litellm

log = logging.getLogger(__name__)

VISION_PROMPT = (
    "Deskripsikan gambar ini untuk asisten AI yang tidak bisa melihatnya. "
    "Salin SEMUA teks yang terlihat apa adanya (struk, dokumen, layar, tanda). "
    "Sebut objek, angka, tabel, dan detail penting. Padat, tanpa basa-basi."
)


class MediaError(RuntimeError):
    pass


@dataclass
class MediaCfg:
    stt_api_base: str = "https://api.groq.com/openai/v1"
    stt_api_key: str = ""
    stt_model: str = "whisper-large-v3-turbo"
    vision_models: list[str] = field(default_factory=lambda: ["gemini/gemini-2.5-flash"])
    file_max_chars: int = 20000
    video_max_minutes: int = 60
    stt_max_mb: int = 24  # batas upload Whisper Groq ~25MB
    ytdlp_cookies: str = ""  # file cookies.txt (Netscape) — YouTube sering minta ini dari IP VPS


def data_url(data: bytes, mime: str) -> str:
    return f"data:{mime};base64,{base64.b64encode(data).decode()}"


class Media:
    def __init__(self, cfg: MediaCfg):
        self.cfg = cfg

    # --- suara ----------------------------------------------------------

    async def transcribe(self, data: bytes, filename: str) -> str:
        if not self.cfg.stt_api_key:
            raise MediaError("STT belum diatur (GROQ_API_KEY kosong).")
        if len(data) > self.cfg.stt_max_mb * 1024 * 1024:
            raise MediaError(f"Audio > {self.cfg.stt_max_mb}MB, terlalu besar untuk STT.")
        async with httpx.AsyncClient(timeout=180) as client:
            resp = await client.post(
                f"{self.cfg.stt_api_base.rstrip('/')}/audio/transcriptions",
                headers={"Authorization": f"Bearer {self.cfg.stt_api_key}"},
                data={"model": self.cfg.stt_model, "response_format": "json"},
                files={"file": (filename, data)},
            )
        if resp.status_code != 200:
            raise MediaError(f"STT HTTP {resp.status_code}: {resp.text[:200]}")
        return (resp.json().get("text") or "").strip()

    # --- gambar ---------------------------------------------------------

    async def describe_image(self, data: bytes, mime: str, hint: str = "") -> str:
        prompt = VISION_PROMPT + (f"\nKonteks dari user: {hint}" if hint else "")
        content = [
            {"type": "text", "text": prompt},
            {"type": "image_url", "image_url": {"url": data_url(data, mime)}},
        ]
        errors = []
        for model in self.cfg.vision_models:
            try:
                resp = await litellm.acompletion(
                    model=model, messages=[{"role": "user", "content": content}], timeout=120
                )
                return (resp.choices[0].message.content or "").strip()
            except Exception as e:
                log.warning("vision %s gagal: %s", model, e)
                errors.append(f"{model}: {str(e)[:120]}")
        raise MediaError("Semua model vision gagal: " + "; ".join(errors))

    # --- file -----------------------------------------------------------

    def extract_file(self, data: bytes, filename: str) -> str:
        ext = Path(filename).suffix.lower()
        if ext == ".pdf":
            text = _pdf(data)
        elif ext == ".docx":
            text = _docx(data)
        elif ext in (".xlsx", ".xlsm"):
            text = _xlsx(data)
        elif ext in TEXT_EXT:
            text = data.decode("utf-8", errors="replace")
        else:
            raise MediaError(f"Tipe file '{ext or filename}' belum didukung.")
        text = text.strip()
        if not text:
            raise MediaError("File tak berisi teks (mungkin hasil scan — kirim sebagai foto).")
        limit = self.cfg.file_max_chars
        if len(text) > limit:
            text = text[:limit] + f"\n\n[…dipotong, total {len(text)} karakter]"
        return text

    # --- video ----------------------------------------------------------

    async def video_transcript(self, url: str) -> str:
        info = await asyncio.to_thread(_yt_info, url, self.cfg.ytdlp_cookies)
        minutes = (info.get("duration") or 0) / 60
        head = (
            f"Judul: {info.get('title')}\nChannel: {info.get('uploader') or info.get('channel')}\n"
            f"Durasi: {minutes:.0f} menit\nURL: {info.get('webpage_url') or url}\n"
        )
        desc = (info.get("description") or "")[:1500]

        sub_url = _pick_subtitle(info)
        if sub_url:
            async with httpx.AsyncClient(timeout=60, follow_redirects=True) as client:
                resp = await client.get(sub_url)
            if resp.status_code == 200:
                text = vtt_to_text(resp.text)
                if text:
                    return self._cut(f"{head}Sumber: subtitle\n\nDeskripsi:\n{desc}\n\nTranskrip:\n{text}")

        if not self.cfg.stt_api_key:
            raise MediaError("Video tak punya subtitle, dan STT belum diatur (GROQ_API_KEY).")
        if minutes > self.cfg.video_max_minutes:
            raise MediaError(
                f"Tak ada subtitle dan video {minutes:.0f} menit > batas "
                f"{self.cfg.video_max_minutes} menit untuk transkrip audio."
            )
        audio, name = await asyncio.to_thread(
            _yt_audio, url, self.cfg.stt_max_mb, self.cfg.ytdlp_cookies
        )
        text = await self.transcribe(audio, name)
        return self._cut(f"{head}Sumber: audio → Whisper\n\nDeskripsi:\n{desc}\n\nTranskrip:\n{text}")

    def _cut(self, text: str) -> str:
        limit = self.cfg.file_max_chars
        return text if len(text) <= limit else text[:limit] + "\n\n[…dipotong]"


TEXT_EXT = {".txt", ".md", ".csv", ".json", ".yaml", ".yml", ".log", ".py", ".js", ".ts",
            ".html", ".xml", ".sql", ".php", ".sh", ".ini", ".toml", ".env.example"}


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    texts = [(p.extract_text() or "").strip() for p in reader.pages]
    if not any(texts):
        return ""
    return "\n".join(f"--- halaman {i} ---\n{t}" for i, t in enumerate(texts, 1))


def _docx(data: bytes) -> str:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"--- sheet {ws.title} ---")
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                out.append(" | ".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


# --- yt-dlp -------------------------------------------------------------

SUB_LANGS = ("id", "en", "en-US", "en-GB")


def _yt_opts(cookies: str, **extra) -> dict:
    opts = {
        "quiet": True,
        "no_warnings": True,
        "noplaylist": True,
        "logger": logging.getLogger("yt_dlp"),  # jangan cetak ke stderr
        **extra,
    }
    if cookies and Path(cookies).exists():
        opts["cookiefile"] = cookies
    return opts


def _yt_error(e: Exception) -> MediaError:
    msg = str(e)
    if "not a bot" in msg or "Sign in" in msg:
        return MediaError(
            "YouTube memblokir IP server ini (minta login). Isi media.ytdlp_cookies "
            "dengan file cookies.txt dari browser, lihat README."
        )
    return MediaError(f"yt-dlp: {msg[:300]}")


def _yt_info(url: str, cookies: str = "") -> dict:
    import yt_dlp

    # Cuma butuh metadata + subtitle; jangan gagal karena pilihan format video.
    opts = _yt_opts(cookies, skip_download=True, ignore_no_formats_error=True)
    with yt_dlp.YoutubeDL(opts) as ydl:
        try:
            return ydl.extract_info(url, download=False)
        except Exception as e:  # extractor yt-dlp kadang rusak dengan error apa saja
            raise _yt_error(e) from e


def _pick_subtitle(info: dict) -> str | None:
    """Utamakan subtitle manual, lalu otomatis; bahasa id → en; format vtt."""
    for key in ("subtitles", "automatic_captions"):
        subs = info.get(key) or {}
        langs = [l for l in SUB_LANGS if l in subs] or [
            l for l in subs if l.split("-")[0] in ("id", "en")
        ]
        # Auto caption YouTube punya banyak terjemahan; ambil bahasa asli kalau ada.
        orig = [l for l in subs if l.endswith("-orig")]
        for lang in orig + langs:
            for f in subs[lang]:
                if f.get("ext") == "vtt":
                    return f["url"]
    return None


def _yt_audio(url: str, max_mb: int, cookies: str = "") -> tuple[bytes, str]:
    import shutil
    import tempfile

    import yt_dlp

    with tempfile.TemporaryDirectory() as tmp:
        opts = _yt_opts(
            cookies,
            # Tanpa ffmpeg: ambil audio asli yang sudah kecil.
            # Situs tanpa audio-saja → video terkecil (Whisper terima mp4/webm).
            format=(
                f"bestaudio[filesize<{max_mb}M]/bestaudio[filesize_approx<{max_mb}M]/"
                f"worstaudio/worst[filesize<{max_mb}M]/worst"
            ),
            outtmpl=f"{tmp}/audio.%(ext)s",
        )
        with yt_dlp.YoutubeDL(opts) as ydl:
            try:
                ydl.extract_info(url, download=True)
            except Exception as e:
                raise _yt_error(e) from e
        files = list(Path(tmp).glob("audio.*"))
        if not files:
            raise MediaError("Audio gagal diunduh.")
        data = files[0].read_bytes()
        if len(data) > max_mb * 1024 * 1024:
            raise MediaError(f"Audio {len(data) // 2**20}MB > {max_mb}MB.")
        return data, files[0].name


_TS = re.compile(r"^\d{2}:\d{2}[:.\d]* --> ")
_TAG = re.compile(r"<[^>]+>")


def vtt_to_text(vtt: str) -> str:
    """WebVTT → teks polos. Buang header, timestamp, tag, dan baris ulang (khas auto-caption)."""
    lines, last = [], None
    for raw in vtt.splitlines():
        line = raw.strip()
        if (
            not line
            or line == "WEBVTT"
            or line.startswith(("Kind:", "Language:", "NOTE", "STYLE"))
            or _TS.match(line)
            or line.isdigit()
        ):
            continue
        line = _TAG.sub("", line).strip()
        if line and line != last:
            lines.append(line)
            last = line
    return " ".join(lines)
