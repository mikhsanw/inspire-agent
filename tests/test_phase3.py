import io

import httpx
import pytest

from inspire.config import ModelCfg
from inspire.llm import strip_images
from inspire.media import Media, MediaCfg, MediaError, _pick_subtitle, vtt_to_text

from test_core import make

VTT = """WEBVTT
Kind: captions
Language: id

00:00:00.000 --> 00:00:02.000
<c>halo</c> semua

00:00:02.000 --> 00:00:04.000
halo semua

00:00:04.000 --> 00:00:06.000
selamat datang
"""


def test_vtt_to_text():
    assert vtt_to_text(VTT) == "halo semua selamat datang"


def test_pick_subtitle_prefers_manual_and_indonesian():
    info = {
        "subtitles": {"en": [{"ext": "vtt", "url": "manual-en"}]},
        "automatic_captions": {"id": [{"ext": "vtt", "url": "auto-id"}]},
    }
    assert _pick_subtitle(info) == "manual-en"
    assert _pick_subtitle({"automatic_captions": info["automatic_captions"]}) == "auto-id"
    assert _pick_subtitle({}) is None


def test_extract_docx_xlsx_text(tmp_path):
    import docx
    from openpyxl import Workbook

    m = Media(MediaCfg(file_max_chars=50))
    d = docx.Document()
    d.add_paragraph("Laporan bulan Oktober")
    buf = io.BytesIO(); d.save(buf)
    assert "Laporan bulan Oktober" in m.extract_file(buf.getvalue(), "a.docx")

    wb = Workbook(); wb.active.append(["nama", "jumlah"]); wb.active.append(["kopi", 3])
    buf = io.BytesIO(); wb.save(buf)
    assert "kopi | 3" in m.extract_file(buf.getvalue(), "b.xlsx")

    long = m.extract_file(b"x" * 100, "c.txt")
    assert long.startswith("x" * 50) and "dipotong" in long

    with pytest.raises(MediaError):
        m.extract_file(b"\x00", "d.exe")
    with pytest.raises(MediaError):
        m.extract_file(b"   ", "e.txt")


def test_strip_images():
    msgs = [{"role": "user", "content": [
        {"type": "text", "text": "ini apa?"},
        {"type": "image_url", "image_url": {"url": "data:..."}},
    ]}]
    out = strip_images(msgs)
    assert isinstance(out[0]["content"], str)
    assert "ini apa?" in out[0]["content"] and "1 gambar" in out[0]["content"]


class FakeMedia:
    async def describe_image(self, data, mime, hint=""):
        return "struk kopi Rp25.000"


async def test_image_non_vision_gets_description(tmp_path):
    agent, llm = make(tmp_path, [{"role": "assistant", "content": "Rp25 ribu"}])
    agent.media = FakeMedia()
    await agent.handle("1", "berapa?", images=[(b"img", "image/jpeg")])
    last = llm.seen[0][-1]["content"]
    assert isinstance(last, str) and "struk kopi Rp25.000" in last and "berapa?" in last
    # deskripsi tersimpan → bisa dicari nanti
    assert agent.store.search("struk", "1")


async def test_image_vision_model_attached_not_stored(tmp_path):
    agent, llm = make(tmp_path, [{"role": "assistant", "content": "kucing"}])
    agent.llm.models["fake"] = ModelCfg(model="fake", vision=True)
    await agent.handle("1", "ini apa?", images=[(b"img", "image/png")])
    last = llm.seen[0][-1]["content"]
    assert isinstance(last, list) and last[1]["image_url"]["url"].startswith("data:image/png;base64,")
    stored = agent.store.context_messages(agent.store.current_session("1"))[0][1]["content"]
    assert isinstance(stored, str) and "gambar terlampir" in stored


async def test_transcribe_posts_to_whisper(monkeypatch):
    seen = {}

    def handler(request: httpx.Request):
        seen["url"] = str(request.url)
        seen["auth"] = request.headers["authorization"]
        seen["body"] = request.content
        return httpx.Response(200, json={"text": " halo dunia "})

    real = httpx.AsyncClient
    monkeypatch.setattr(httpx, "AsyncClient",
                        lambda **kw: real(transport=httpx.MockTransport(handler), **kw))
    m = Media(MediaCfg(stt_api_key="k"))
    assert await m.transcribe(b"OggS", "voice.ogg") == "halo dunia"
    assert seen["url"].endswith("/audio/transcriptions") and seen["auth"] == "Bearer k"
    assert b"whisper-large-v3-turbo" in seen["body"] and b"voice.ogg" in seen["body"]

    with pytest.raises(MediaError):
        await Media(MediaCfg(stt_api_key="")).transcribe(b"x", "a.ogg")
